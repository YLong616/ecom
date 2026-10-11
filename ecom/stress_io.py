"""压力模式：AT list 模版 / 导入 / 导出（CSV / TXT / XLSX）。

跨平台约定：
- CSV/TXT 用 ``utf-8-sig``（带 BOM）写，Windows Excel 双击不乱码，Linux 亦可正常读；
  读兼容带 BOM 的 utf-8-sig 与无 BOM 的 utf-8，满足双系统来回导入。
- ``.xlsx``（Excel 真文件，需 openpyxl，pip install openpyxl）用于「导出模版」：
  表头 + 示例行 + 下拉框，让 send_type / send_suffix / match_mode / enabled / hex / enter
  直接在 Excel 里下拉选择，避免手填出错。
- 数据导入兼容 .csv / .txt / .xlsx；数据导出可选 .csv / .txt（.xlsx 亦可）。

列（已去掉展示用 index，共 9 列）：enabled / at / send_type / send_suffix /
hex / enter / interval_ms / expected / match_mode。
兼容旧 10 列（含 index 于首列）：读入时依据表头首列文字自动识别并偏移。
"""

from __future__ import annotations

import csv
import io
import os

# 列头（写死英文码值，跨语言/跨平台稳定；匹配模式列存 contains/regex/exact）
HEADERS = [
    "enabled", "at", "send_type", "send_suffix",
    "hex", "enter", "interval_ms", "expected", "match_mode",
]

# 导出模版示例（预期结果可为多个，用 | 分隔）
TEMPLATE_ROWS = [
    {
        "enabled": True, "txt": "AT+CPIN?", "send_type": "ASCII",
        "send_suffix": "CR-LF", "hex": False, "enter": False,
        "interval_ms": 0, "expected": "OK", "match_mode": "contains",
    },
    {
        "enabled": True, "txt": "AT+CIMI", "send_type": "ASCII",
        "send_suffix": "CR-LF", "hex": False, "enter": False,
        "interval_ms": 0, "expected": "", "match_mode": "contains",
    },
    {
        "enabled": True, "txt": "AT+CCID", "send_type": "ASCII",
        "send_suffix": "CR-LF", "hex": False, "enter": False,
        "interval_ms": 0, "expected": "", "match_mode": "contains",
    },
]

# 运行时使用的行字段集合
FIELD_KEYS = [
    "enabled", "txt", "send_type", "send_suffix",
    "hex", "enter", "interval_ms", "expected", "match_mode",
]

# 枚举值（下拉框 + 导入归一化共用）
SEND_TYPES = ("ASCII", "HEX", "DEC", "BIN")
SEND_SUFFIXES = ("None", "CR", "LF", "CR-LF", "SPACE", "STX-ETX", "NULL")
MATCH_MODES = ("contains", "regex", "exact")


def _norm_bool(v) -> bool:
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return bool(v)
    return str(v).strip().lower() in ("1", "true", "yes", "y", "on")


def _norm_int(v, default: int = 0) -> int:
    try:
        return int(str(v).strip())
    except (ValueError, TypeError):
        return default


def _offset(header) -> int:
    """旧 10 列（index 在首列）→ 偏移 1；新 9 列 → 偏移 0。"""
    if header and str(header[0]).strip().lower() == "index":
        return 1
    return 0


def _row_to_list(row: dict) -> list:
    return [
        "1" if _norm_bool(row.get("enabled", True)) else "0",
        str(row.get("txt", "")),
        str(row.get("send_type", "ASCII")),
        str(row.get("send_suffix", "None")),
        "1" if _norm_bool(row.get("hex", False)) else "0",
        "1" if _norm_bool(row.get("enter", False)) else "0",
        str(_norm_int(row.get("interval_ms"), 0)),
        str(row.get("expected", "")),
        str(row.get("match_mode", "contains")),
    ]


def _row_from_list(fields: list, offset: int = 0) -> dict:
    """按列位置解析一行；缺列/多列容忍。offset 用于兼容旧 10 列（index 占首列）。"""
    def get(i, default=""):
        j = i + offset
        return fields[j].strip() if j < len(fields) else default

    row = {
        "enabled": _norm_bool(get(0, "1")),
        "txt": get(1, ""),
        "send_type": get(2, "ASCII"),
        "send_suffix": get(3, "None"),
        "hex": _norm_bool(get(4, "0")),
        "enter": _norm_bool(get(5, "0")),
        "interval_ms": _norm_int(get(6, "0"), 0),
        "expected": get(7, ""),
        "match_mode": get(8, "contains"),
    }
    if row["send_type"] not in SEND_TYPES:
        row["send_type"] = "ASCII"
    if row["send_suffix"] not in SEND_SUFFIXES:
        row["send_suffix"] = "None"
    if row["match_mode"] not in MATCH_MODES:
        row["match_mode"] = "contains"
    return row


# ---------- CSV / TXT（内容一样，逗号分隔，按扩展名路由） ----------
def _write_csv(path: str, rows: list[dict]) -> bool:
    try:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(HEADERS)
            for row in rows:
                w.writerow(_row_to_list(row))
        return True
    except OSError:
        return False


def _read_csv(path: str) -> list[dict]:
    rows = []
    try:
        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            data = f.read()
        reader = csv.reader(io.StringIO(data))
        header = None
        for field in reader:
            if header is None:
                header = field
                continue
            if not field or not any(c.strip() for c in field):
                continue
            rows.append(_row_from_list(field, _offset(header)))
    except (OSError, UnicodeDecodeError):
        try:
            with open(path, "r", newline="", encoding="utf-8") as f:
                reader = csv.reader(f)
                header = next(reader, None)
                for field in reader:
                    if not field or not any(c.strip() for c in field):
                        continue
                    rows.append(_row_from_list(field, _offset(header)))
        except (OSError, UnicodeDecodeError):
            return []
    return rows


# ---------- XLSX（需 openpyxl） ----------
def _add_dropdowns(ws) -> None:
    """给枚举列加 Excel 验证下拉框（应用到 2..2000 行，覆盖示例与后续填写）。"""
    from openpyxl.worksheet.datavalidation import DataValidation

    def dv(opts):
        d = DataValidation(type="list", formula1='"' + ",".join(opts) + '"',
                           allow_blank=True)
        return d

    def apply(d, col):
        ws.add_data_validation(d)
        d.add("%s2:%s2000" % (col, col))

    apply(dv(["0", "1"]), "A")                        # enabled
    apply(dv(list(SEND_TYPES)), "C")                  # send_type
    apply(dv(list(SEND_SUFFIXES)), "D")               # send_suffix
    apply(dv(["0", "1"]), "E")                        # hex
    apply(dv(["0", "1"]), "F")                        # enter
    apply(dv(list(MATCH_MODES)), "I")                 # match_mode


def _write_xlsx(path: str, rows: list[dict], with_dropdowns: bool) -> bool:
    try:
        from openpyxl import Workbook
    except ImportError:
        return False
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Stress"
        ws.append(HEADERS)
        for row in rows:
            ws.append(["" if v is None else str(v) for v in _row_to_list(row)])
        if with_dropdowns:
            _add_dropdowns(ws)
        wb.save(path)
        return True
    except Exception:
        return False


def _read_xlsx(path: str) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    rows = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = ["" if v is None else str(v) for v in (next(it, None) or [])]
        off = _offset(header)
        for cells in it:
            field = ["" if v is None else str(v) for v in cells]
            if not field or not any(c.strip() for c in field):
                continue
            rows.append(_row_from_list(field, off))
        wb.close()
    except Exception:
        return []
    return rows


# ---------- 对外 API ----------
def export_template(path: str) -> bool:
    """导出带示例行的模版。.xlsx → 带下拉框；.csv/.txt → 逗号分隔纯文本。"""
    if str(path).lower().endswith(".xlsx"):
        return _write_xlsx(path, TEMPLATE_ROWS, with_dropdowns=True)
    return _write_csv(path, TEMPLATE_ROWS)


def export_data(path: str, rows: list[dict]) -> bool:
    """导出当前 AT list 全量数据。扩展名决定格式：.csv/.txt → CSV，.xlsx → XLSX。"""
    if str(path).lower().endswith(".xlsx"):
        return _write_xlsx(path, rows, with_dropdowns=False)
    return _write_csv(path, rows)


def import_data(path: str) -> list[dict]:
    """导入本地 CSV/TXT/XLSX，返回行 dict 列表（可为空）。"""
    if str(path).lower().endswith(".xlsx"):
        return _read_xlsx(path)
    return _read_csv(path)
