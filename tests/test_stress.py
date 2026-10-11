# -*- coding: utf-8 -*-
"""压力模式纯逻辑单测（无 Tk，mock 串口可另测）。py -3.8 运行。"""
import io
import os
import sys
import tempfile
import unittest
from datetime import datetime

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from ecom import config, i18n, stress_io, stress_model  # noqa: E402
from ecom.stress_model import (MATCH_CONTAINS, MATCH_EXACT, MATCH_REGEX, TARGET_AT,
                               TARGET_DOS, CheckRule, StressRow,
                               compute_end_condition, format_duration, format_window,
                               match_expected, parse_expected, run_check_rules)


class TestExpected(unittest.TestCase):
    def test_parse_expected_newline_and_pipe(self):
        self.assertEqual(parse_expected("OK|ERROR\nFAIL"), ["OK", "ERROR", "FAIL"])

    def test_parse_expected_blank(self):
        self.assertEqual(parse_expected("   "), [])

    def test_match_contains(self):
        self.assertTrue(match_expected("AT+CPIN? response OK", "OK", MATCH_CONTAINS))
        self.assertFalse(match_expected("ERROR response", "OK", MATCH_CONTAINS))

    def test_match_exact(self):
        self.assertTrue(match_expected("OK", "OK", MATCH_EXACT))
        # 全等比较会 strip 首尾空白，视为一致
        self.assertTrue(match_expected("OK\r\n", "OK", MATCH_EXACT))
        self.assertFalse(match_expected("OK2", "OK", MATCH_EXACT))

    def test_match_regex(self):
        self.assertTrue(match_expected("+QENG: 1,2,3", r"\+\w+\s*:", MATCH_REGEX))
        self.assertFalse(match_expected("hello", r"^\d+$", MATCH_REGEX))

    def test_match_any_of_multiple(self):
        self.assertTrue(match_expected("ERROR here", "OK|ERROR", MATCH_CONTAINS))

    def test_match_no_expected_passes(self):
        self.assertTrue(match_expected("anything", "", MATCH_CONTAINS))

    def test_invalid_regex_fallback(self):
        # 非法正则退化为包含匹配：能命中子串则 True
        self.assertTrue(match_expected("text [ here", "[", MATCH_REGEX))
        self.assertFalse(match_expected("some text", "[", MATCH_REGEX))


class TestRules(unittest.TestCase):
    def test_run_rules_hits(self):
        rules = [CheckRule(name="RDY", keyword="RDY", record_to=TARGET_DOS,
                           as_exception=True, enabled=True)]
        hits = run_check_rules("some RDY line", rules)
        self.assertEqual(len(hits), 1)

    def test_run_rules_disabled_skipped(self):
        rules = [CheckRule(name="RDY", keyword="RDY", record_to=TARGET_DOS,
                           as_exception=True, enabled=False)]
        self.assertEqual(run_check_rules("RDY", rules), [])


class TestRows(unittest.TestCase):
    def test_effective_type_suffix(self):
        r = StressRow(txt="AT+CPIN?", send_type="ASCII", send_suffix="None",
                      hex=True, enter=True)
        self.assertEqual(r.effective_type, "HEX")
        self.assertEqual(r.effective_suffix, "CR-LF")

    def test_to_from_dict_roundtrip(self):
        r = StressRow(txt="AT", send_type="HEX", send_suffix="CR-LF", hex=False,
                      enter=True, interval_ms=100, expected="OK|ERROR",
                      match_mode=MATCH_REGEX, enabled=True)
        d = r.to_dict()
        r2 = StressRow.from_dict(d)
        self.assertEqual(r2.txt, "AT")
        self.assertEqual(r2.match_mode, MATCH_REGEX)
        self.assertEqual(r2.interval_ms, 100)
        self.assertTrue(r2.enter)

    def test_invalid_match_mode_normalized(self):
        r = StressRow(match_mode="bogus")
        self.assertEqual(r.match_mode, MATCH_CONTAINS)


class TestStats(unittest.TestCase):
    def test_format_duration(self):
        self.assertEqual(format_duration(0), "0.0H")
        self.assertEqual(format_duration(1800), "0.5H")

    def test_format_window(self):
        a = datetime(2026, 8, 25, 12, 0, 0, 0)
        b = datetime(2026, 8, 25, 12, 5, 0, 0)
        s = format_window(a, b)
        self.assertIn("2026-08-25 12:00:00.000000", s)
        self.assertIn("~", s)

    def test_end_condition_count(self):
        self.assertTrue(compute_end_condition(1, "count", 1, 0.0, 0.0))
        self.assertFalse(compute_end_condition(0, "count", 1, 0.0, 0.0))

    def test_end_condition_time(self):
        self.assertTrue(compute_end_condition(1, "time", 1, 60.0, 60.0))
        self.assertFalse(compute_end_condition(1, "time", 1, 30.0, 60.0))


class TestIo(unittest.TestCase):
    def _tmp(self, suffix=".csv"):
        fd, p = tempfile.mkstemp(suffix=suffix)
        os.close(fd)
        self.addCleanup(lambda: os.path.exists(p) and os.remove(p))
        return p

    def test_template_roundtrip(self):
        p = self._tmp()
        self.assertTrue(stress_io.export_template(p))
        rows = stress_io.import_data(p)
        self.assertTrue(len(rows) >= 1)
        # 第一个示例行应含预期结果
        self.assertIn("OK", rows[0]["expected"])

    def test_export_import_roundtrip_fields(self):
        p = self._tmp()
        src = [
            {"enabled": True, "txt": "AT+CPIN?", "send_type": "ASCII",
             "send_suffix": "CR-LF", "hex": False, "enter": False,
             "interval_ms": 50, "expected": "OK|ERROR", "match_mode": "contains"},
            {"enabled": False, "txt": "AT+QENG=?", "send_type": "HEX",
             "send_suffix": "None", "hex": True, "enter": True,
             "interval_ms": 0, "expected": "RDY", "match_mode": "regex"},
        ]
        self.assertTrue(stress_io.export_data(p, src))
        rows = stress_io.import_data(p)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["txt"], "AT+CPIN?")
        self.assertEqual(rows[0]["match_mode"], "contains")
        self.assertEqual(rows[0]["interval_ms"], 50)
        self.assertTrue(rows[1]["hex"])
        self.assertTrue(rows[1]["enter"])
        self.assertFalse(rows[1]["enabled"])
        self.assertEqual(rows[1]["match_mode"], "regex")

    def test_import_bom_and_nobom(self):
        p = self._tmp()
        data = "enabled,at,send_type,send_suffix,hex,enter,interval_ms,expected,match_mode\n" \
               "1,AT+CPIN?,ASCII,CR-LF,0,0,0,OK,contains\n"
        with open(p, "w", encoding="utf-8-sig") as f:
            f.write(data)
        rows = stress_io.import_data(p)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["txt"], "AT+CPIN?")

    def test_data_export_txt_roundtrip(self):
        p = self._tmp(".txt")
        src = [{"enabled": True, "txt": "AT+CPIN?", "send_type": "HEX",
                "send_suffix": "CR-LF", "hex": True, "enter": False,
                "interval_ms": 30, "expected": "OK", "match_mode": "exact"}]
        self.assertTrue(stress_io.export_data(p, src))
        rows = stress_io.import_data(p)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["txt"], "AT+CPIN?")
        self.assertEqual(rows[0]["send_type"], "HEX")
        self.assertTrue(rows[0]["hex"])

    def test_xlsx_template_has_dropdowns(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl 不可用")
        p = self._tmp(".xlsx")
        self.assertTrue(stress_io.export_template(p))
        wb = openpyxl.load_workbook(p)
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "enabled")   # 无 index 列
        self.assertGreaterEqual(len(ws["D"]), 3)                      # 至少 2 行示例 + 表头
        dv_cols = set()
        for dv in ws.data_validations.dataValidation:
            for rng in dv.sqref.ranges:
                dv_cols.add(str(rng)[0])
        self.assertIn("C", dv_cols)     # send_type
        self.assertIn("D", dv_cols)     # send_suffix
        self.assertIn("I", dv_cols)     # match_mode
        wb.close()

    def test_import_xlsx_data_roundtrip(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl 不可用")
        p = self._tmp(".xlsx")
        src = [{"enabled": True, "txt": "AT+QENG=?", "send_type": "ASCII",
                "send_suffix": "None", "hex": False, "enter": True,
                "interval_ms": 100, "expected": "OK|ERROR", "match_mode": "contains"}]
        self.assertTrue(stress_io.export_data(p, src))
        rows = stress_io.import_data(p)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["txt"], "AT+QENG=?")
        self.assertEqual(rows[0]["expected"], "OK|ERROR")
        self.assertTrue(rows[0]["enter"])


class TestConfigI18nApp(unittest.TestCase):
    def test_config_stress_section(self):
        c = config.Config()
        self.assertIn("stress", c.data)
        self.assertIn("end_mode", c.data["stress"])

    def test_i18n_stress_keys(self):
        for key in ["mode_at", "mode_stress", "stress_start", "stress_stop",
                    "stress_at_list", "stress_error"]:
            pass
        self.assertEqual(i18n.get("zh", "mode_stress"), "压力模式")
        self.assertEqual(i18n.get("en", "mode_stress"), "Stress Mode")
        self.assertEqual(i18n.get("zh", "stress_no_rows"), "请先勾选至少一行")

    def test_app_mode_attrs(self):
        # app 是 GUI，不能实例化；这里只做源码级断言，确保补丁没破坏关键锚点
        app_path = os.path.join(os.path.dirname(__file__), "..", "ecom", "app.py")
        with open(app_path, encoding="utf-8", errors="replace") as fh:
            app_src = fh.read()
        self.assertIn("self._mode = \"at\"", app_src)
        self.assertIn("def _switch_mode", app_src)
        self.assertIn("self.at_view = body", app_src)
        self.assertIn("from .stress_panel import StressPanel", app_src)


if __name__ == "__main__":
    unittest.main(verbosity=2)
